"""Build the hackathon submission report (.docx) with python-docx. Numbers come from data/state.json and
data/field_notes.json at build time so the document never carries stale counts.
`uv run --with python-docx python scripts/build_submission_docx.py` → docs/[NVIDIA 해커톤_AI Pioneer_DELPHi].docx"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "[NVIDIA 해커톤_AI Pioneer_DELPHi].docx"
NAVY, GREY = RGBColor(0x16, 0x26, 0x61), RGBColor(0x55, 0x55, 0x55)
CONSOLE = "https://delphi-console-production-8ae8.up.railway.app"
APP = "https://delphi-web-production-52d6.up.railway.app"
REPO = "https://github.com/coldtype-08/delphi-evidence-loop"
SIGNAL_KO = {"OFF_LABEL_DEMAND": "허가 밖 수요", "OFF_LABEL_USE": "허가 밖 사용 경험", "REPURPOSING": "다른 용도",
             "UNMET_NEED": "충족되지 않은 필요", "DOSING": "용량과 제형", "SAFETY_TOLERABILITY": "안전성과 내약성"}
RECO_KO = {"PROCEED_TO_EXPERT_REVIEW": "전문조직 검토 권고", "HOLD": "보류 권고", "DROP": "기각 권고"}
DEC_KO = {"PROCEED_TO_EXPERT_REVIEW": "전문조직 검토로 결정", "HOLD": "보류로 결정", "DROP": "기각으로 결정"}


# ── data ──────────────────────────────────────────────────────────────────────

def facts() -> dict:
    st = json.loads((ROOT / "data" / "state.json").read_text())
    notes = json.loads((ROOT / "data" / "field_notes.json").read_text())
    contract = json.loads((ROOT / "data" / "contract.json").read_text())
    claims = st["claims"]
    hyps = st["hypotheses"]
    rows = []
    for h in hyps:
        s = st["screens"].get(h["id"])
        m = st["board"].get(h["id"])
        dec = (m or {}).get("decision")
        rows.append({
            "id": h["id"], "pair": f'{h["segment"]} × {SIGNAL_KO.get(h["signal_type"], h["signal_type"])}',
            "field": f'{h["field"]["mentions"]}회 / {h["field"]["hcps"]}인',
            "evidence": (f'지지 {s["totals"]["SUPPORTS"]} · 반대 {s["totals"]["CONTRADICTS"]} · 중립 {s["totals"]["NEUTRAL"]}' if s else "미실시"),
            "label": "허가 범위 밖" if h["label_status"] == "DEVELOPMENT" else "허가 범위 안",
            "board": (f'{len(m["transcript"])}턴, 지지 {m["tally"]["counts"]["SUPPORT"]} 보류 {m["tally"]["counts"]["HOLD"]} 반대 {m["tally"]["counts"]["OPPOSE"]}, {RECO_KO.get(m["recommendation"], m["recommendation"])}'
                      if m and "transcript" in m else "미실시"),
            "decision": (f'{DEC_KO.get(dec["accepted"], dec["accepted"])} ({dec["by"]}), 후속 질문 {len([a for a in st["actions"] if a["hypothesis_id"] == h["id"]])}건'
                         if dec else ("판정 대기" if m else "")),
            "statement": h["statement_ko"],
        })
    return {
        "notes": len(notes), "hcps": len({n["hcp_ref"] for n in notes}), "period": f'{min(n["date"] for n in notes)} ~ {max(n["date"] for n in notes)}',
        "claims": len(claims), "verified": sum(c["verified"] for c in claims), "ae": len(st["safety_queue"]),
        "hyps": len(hyps), "screened": len(st["screens"]), "boards": len([m for m in st["board"].values() if "transcript" in m]),
        "decided": len([m for m in st["board"].values() if m.get("decision")]), "actions": len(st["actions"]),
        "threshold": contract["threshold"], "segments": contract["segments"], "rows": rows, "state": st,
    }


# ── docx helpers ──────────────────────────────────────────────────────────────

def hyperlink(paragraph, url: str, text: str | None = None):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "162661"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text or url; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run); paragraph._p.append(link)


def para(doc, text="", *, bold=False, size=None, color=None, after=6, align=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h


def table(doc, header, rows, widths=None, size=9.5):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


WITH_PICTURES = False   # 제출본은 글로만 — 캡처는 실행 시점마다 달라 본문 숫자와 어긋난다


def picture(doc, name, caption):
    path = DOCS / "shots" / name
    if not WITH_PICTURES or not path.exists():
        return
    doc.add_picture(str(path), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=9, color=GREY, after=10)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet"); p.paragraph_format.space_after = Pt(2)


# ── document ──────────────────────────────────────────────────────────────────

def build():
    f = facts()
    thr = f["threshold"]
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Apple SD Gothic Neo"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Apple SD Gothic Neo")
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)

    # 표지
    para(doc, "AI Pioneer", bold=True, size=22, color=NAVY, after=2)
    para(doc, "DELPHi: 의료진 면담 기록에서 시장 신호를 세고, 공개 근거로 확인하고, 사람이 결정하는 시스템", size=12, color=GREY, after=0)
    para(doc, "Korea Agentic AI Hackathon 온라인 예선 제출 · 2026년 9월 28일", size=11, color=GREY, after=12)

    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, url, label in [("배포 콘솔", CONSOLE, CONSOLE), ("API와 소개 페이지", APP, APP), ("GitHub", REPO, REPO),
                          ("실행 기록 전문", f"{REPO}/blob/main/docs/demo_run.txt", "docs/demo_run.txt")]:
        c = t.add_row().cells; c[0].text = ""; c[1].text = ""
        rr = c[0].paragraphs[0].add_run(k); rr.bold = True; rr.font.size = Pt(9.5)
        hyperlink(c[1].paragraphs[0], url, label)
        c[0].width = Inches(1.6); c[1].width = Inches(5.1)
    doc.add_paragraph()

    # 1. 요약
    heading(doc, "1. 요약")
    para(doc, f"제약회사 의학부가 의료진을 만나고 남기는 면담 기록을 입력으로 받아, 환자군과 신호 유형별로 세고, 반복이 일정 수를 넘은 조합을 가설로 만들어 "
              f"공개 문헌과 임상시험 등록, 허가사항과 대조한 뒤, 임원 역할 에이전트들의 심의를 거쳐 사람이 결정하는 시스템이다. "
              f"이번 데모는 특허가 만료된 메트포르민을 대상으로 가상 의료진 {f['hcps']}인의 면담 {f['notes']}건을 합성해 돌렸다. "
              f"추출된 발언 {f['claims']}건 중 {f['verified']}건이 원문 위치까지 확인됐고, 유해사례 후보 {f['ae']}건은 별도 경로로 분리됐다. "
              f"문턱(반복 {thr['min_mentions']}회, 의료진 {thr['min_hcps']}인)을 넘은 가설 {f['hyps']}건에 대해 근거 조사 {f['screened']}건, AI Board 심의 {f['boards']}건, 사람의 결정 {f['decided']}건이 이루어졌고, "
              f"결정에서 나온 후속 질문 {f['actions']}건이 다음 면담 체크리스트에 올라가 있다.")
    para(doc, "이 문서의 숫자는 배포된 상태 파일에서 문서를 만들 때 읽어 온 값이다. 화면의 숫자와 다르면 화면이 맞다.", size=9.5, color=GREY)

    # 2. 문제
    heading(doc, "2. 해결하려는 문제")
    heading(doc, "2.1 기록은 쌓이는데 셀 수가 없다", 2)
    para(doc, "의학부 담당자는 의료진을 만난 뒤 자유 형식으로 기록을 남긴다. 그 안에는 허가 범위 밖 환자에게 쓰고 싶다는 말, 실제로 써 본 경험, "
              "다른 용도로 쓰고 있다는 이야기가 섞여 있다. 문서마다 형식과 표현이 달라 같은 이야기가 몇 명에게서 몇 번 나왔는지 알 수 없고, "
              "필요할 때 사람이 문서를 다시 열어 읽고 정리한다. 문서 한 건은 일화지만 열 건이 모이면 회사가 검토해야 할 신호가 된다. "
              "얼마나 쌓였는지 자체가 보이지 않으니 가설은 담당자의 기억에 기댄다. 업계 설문에서도 의학부 팀의 91%가 수집한 데이터 대부분을 인사이트로 쓰지 못한다고 답했다(Within3와 Reuters 설문, 업계 자료).")
    heading(doc, "2.2 AI 요약은 근거가 되지 못한다", 2)
    para(doc, "의료와 규제 영역에서는 누가 언제 무슨 말을 했는지 원문까지 되짚을 수 있어야 판단에 올릴 수 있다. 언어모델에 요약을 맡기면 "
              "횟수를 지어내거나 근거 없는 문장이 섞이고, 그 결과는 어떤 회의에도 가져갈 수 없다. 요약이 아니라 원문 위치가 붙은 근거가 필요했다.")
    heading(doc, "2.3 허가 범위 밖 신호는 다루는 규칙이 따로 있다", 2)
    para(doc, "미승인 적응증이나 환자군에 대한 수요는 전문조직이 검토할 대상이지 상업 활동의 재료가 아니다. 또 「써 봤다」와 「쓰고 싶은데 막혔다」는 "
              "규제상 다른 신호라서 한 칸에 섞으면 안 되고, 개별 환자의 유해사례로 읽히는 발언은 처음부터 다른 경로로 가야 한다. "
              "이 규칙을 문서로 약속하는 대신 코드가 지키게 하는 것이 설계의 출발점이었다.")

    # 3. 해결 방안
    heading(doc, "3. 해결 방안")
    heading(doc, "3.1 처리 순서 6단계", 2)
    table(doc, ["단계", "주체", "하는 일"], [
        ["① 추출", "Nemotron과 코드", "면담 기록에서 환자군과 신호 유형에 해당하는 문장을 골라 원문 그대로 인용한다. 코드가 인용을 원문에서 찾아 문자 위치를 붙이고, 못 찾으면 버린다. 유해사례로 읽히는 문장은 safety 큐로 보낸다."],
        ["② 계수와 문턱", "코드", f"언급 횟수와 독립 의료진 수를 센다. 반복 {thr['min_mentions']}회, 의료진 {thr['min_hcps']}인을 넘은 (환자군 × 신호 유형) 조합만 가설이 된다. 가설 문장과 검색식은 모델이 쓰지만 숫자는 쓰지 않는다."],
        ["③ 근거 조사", "판독 에이전트 3", "PubMed, ClinicalTrials.gov, FDA 허가사항을 읽고 기록마다 지지, 반대, 중립을 원문 인용과 함께 표시한다. 무작위 대조시험과 3상, 메타분석을 먼저 읽는다. FAERS 보고 건수와 Medicare Part D 청구 건수는 API 값을 그대로 쓴다."],
        ["④ 서명", "사람", "근거를 직접 읽었다는 서명을 이름으로 남긴다. 서명이 없으면 심의가 열리지 않는다. 반대 근거가 많아도 상정할 수 있다."],
        ["⑤ 심의", "간사 1과 임원 7", "개회, 모두발언, 쟁점 토론, 최종 입장 순으로 회의를 연다. 발언의 인용은 근거 목록에 있는 ID만 인정한다. 최종 입장은 코드가 확신도로 가중해 집계하고 그 결과가 권고가 된다. 간사는 회의록만 쓴다."],
        ["⑥ 판정과 실행", "사람", "의장이 추진, 조건부 추진, 보류, 기각 중 하나를 판정한다. 회의록의 후속 질문이 다음 면담 체크리스트로 내려간다."],
    ], [1.0, 1.2, 4.5])
    heading(doc, "3.2 사람이 개입하는 자리 두 곳", 2)
    para(doc, "추출된 발언을 한 건씩 승인하지 않는다. 사람은 가설 단위로 근거를 검토해 서명하고, 심의가 끝나면 판정한다. "
              "두 자리 모두 이름과 시각이 기록에 남는다. 검토량은 데이터 양이 아니라 실제로 판단에 쓰이는 양에 비례한다.")
    heading(doc, "3.3 코드가 지키는 규칙 5개", 2)
    table(doc, ["규칙", "구현"], [
        ["숫자는 모델이 만들지 않는다", "언급 수, 의료진 수, 지지와 반대 건수, 시험 수, 청구 건수는 전부 코드가 센다. 심의의 권고도 집계 결과다."],
        ["인용에는 원문 위치가 있다", "모델이 준 인용문을 원문에서 찾지 못하면 「버림」으로 표시하고 세지 않는다. 문장 사이를 건너뛰어 이어 붙인 인용도 버린다."],
        ["관문은 둘이고 이름이 남는다", "서명 없이 심의를 열 수 없고, 판정 없이 실행 항목이 생기지 않는다."],
        ["허가 범위 밖 가설은 전문조직 검토로만 간다", "심의에서 나온 상업 액션(프로모션, 영업 목표 등)은 코드가 걸러 기록에 남긴다."],
        ["유해사례 후보는 분석에 섞이지 않는다", "추출 단계에서 별도 큐로 분리하고, 집계와 가설 생성에서 제외한다."],
    ], [2.0, 4.7])

    # 4. 시스템 구성
    heading(doc, "4. 시스템 구성")
    heading(doc, "4.1 화면 7개", 2)
    para(doc, "홈(현황과 신호 지도), 신호의 여정(파이프라인 보드), 신호와 가설(가설 목록과 생성), 다중 에이전트 검증(근거 조사 워크벤치), "
              "심의(안건 대장과 AI Board 회의실), 안전(유해사례 후보), 실행 기록(모델 호출 이력과 비용). 회의실은 발언이 생기는 순서대로 화면에 올라오고, 끝난 회의는 처음부터 다시 재생할 수 있다.")
    picture(doc, "console_home.png", "홈 화면. 면담 코퍼스 규모, 허가 범위 밖과 안의 언급 건수, 환자군별 추세, 신호 유형 분포, 환자군과 신호 유형의 교차 집계를 한 화면에 둔다.")
    picture(doc, "journey.png", "신호의 여정. 현장, SENSE, SCREEN, AI BOARD 네 레인에 단계별 건수를 놓고, 코드가 세는 단계와 AI가 판단하는 단계, 사람이 개입하는 단계를 구분해 표시한다.")
    heading(doc, "4.2 NVIDIA 스택과 에이전트 구성", 2)
    table(doc, ["구분", "내용"], [
        ["추론 모델", "NVIDIA Nemotron 3 Ultra (nvidia/nemotron-3-ultra-550b-a55b). NIM의 OpenAI 호환 API로 호출한다. 한국어를 공식 지원해 면담 기록을 번역 없이 읽는다."],
        ["호출 방식", "함수 호출로 JSON 스키마를 강제하고 jsonschema로 검증한다. 회의록을 쓰는 단계에서만 reasoning 모드를 켠다. 모든 호출은 래퍼 한 곳을 지나며 응답을 캐시하고 실행 기록(모델, 토큰, 캐시 적중)을 남긴다. 속도 제한(429)에는 지수 백오프로 대응한다."],
        ["에이전트", "판독 에이전트 3종(PubMed, ClinicalTrials.gov, FDA 라벨), 간사 1, 임원 7(CMO, 규제, 약물감시, 임상개발, 재무, 상업, CEO). 임원은 병렬로 발언하고 간사가 지명해 토론한다."],
        ["스킬과 샌드박스", "NVIDIA Agent Skills 규격의 evidence-loop 스킬(skills/evidence-loop/SKILL.md)과 OpenShell 정책용 허용 호스트 5개 명세(sandbox/EGRESS.md)를 함께 제출한다. 샌드박스 실행은 본선에서 진행한다."],
        ["앱", "콘솔 Next.js 16, React 19, Tailwind 4. 백엔드 Python 3.12, FastAPI, uv. 상태는 JSON 파일 하나가 정본이다."],
        ["근거원", "PubMed E-utilities, ClinicalTrials.gov API v2, openFDA(허가사항과 FAERS), CMS Medicare Part D Spending. 응답은 디스크에 캐시한다."],
        ["배포", "Railway 2개 서비스. 콘솔은 Nixpacks 빌드 뒤 응답을 버퍼링하는 Node 서버로 서비스하고, API는 Docker로 올린다."],
    ], [1.4, 5.3])

    # 5. 데모 결과
    heading(doc, "5. 데모 결과")
    heading(doc, "5.1 데이터", 2)
    para(doc, f"대상 약은 메트포르민이다. 특허가 만료됐고 특정 회사 소유가 아니며 공개 근거가 많다(PubMed 제목 검색 18,054편, ClinicalTrials.gov 3,120건, FAERS 보고 440,270건, Medicare Part D 2024년 청구 3,437만 건, 2026년 9월 28일 조회). "
              f"면담 기록은 가상 의료진 {f['hcps']}인의 {f['notes']}건({f['period']})을 합성했다. 문장은 공개된 사실에 맞춰 썼다. "
              "예를 들어 유방암 보조요법에 대해서는 3,649명 규모의 캐나다 주도 3상 시험이 무효로 끝났다는 사실을, 신기능 저하 환자에 대해서는 eGFR 30 미만 금기와 30~45 신규 투여 금지, 젖산산증 박스 경고를, "
              "다낭성난소증후군에 대해서는 레트로졸이나 클로미펜과의 병용이 허가 밖이지만 널리 쓰인다는 사실을 반영했다. 실제 환자, 의료진, 기관 정보는 없다.")
    heading(doc, f"5.2 가설 {f['hyps']}건의 처리 결과", 2)
    table(doc, ["가설", "환자군 × 신호", "현장", "외부 근거", "허가", "AI Board", "결정"],
          [[r["id"], r["pair"], r["field"], r["evidence"], r["label"], r["board"], r["decision"]] for r in f["rows"]],
          [0.6, 1.5, 0.7, 1.4, 0.7, 1.2, 1.1], size=8.5)
    para(doc, "현장 칸은 언급 횟수와 독립 의료진 수, 외부 근거 칸은 판독 에이전트가 표시한 건수(코드 집계)다. 인용을 원문에서 찾지 못해 버린 건수는 화면에 따로 표시된다.", size=9.5, color=GREY)
    picture(doc, "hyp_detail.png", "가설 상세. 현장 발언(원문 위치), 검색식, 외부 근거 지표, 판정별 근거 표를 한 화면에 둔다.")
    picture(doc, "boardroom_top.png", "AI Board 회의실. 회의 정보, 참석자 레일, 간사의 소집 질의, 임원 발언 순서로 재생된다.")
    picture(doc, "boardroom_verdict.png", "회의 끝의 의장 판정 패널. 최종 입장 집계와 권고를 보여 주고, 추진, 조건부 추진, 보류, 기각 중 하나를 의장 서명과 함께 받는다.")
    picture(doc, "checklist.png", "다음 면담 체크리스트. 판정에서 채택된 후속 질문이 항목으로 내려간다.")

    # 6. 심사 기준
    heading(doc, "6. 심사 기준별 대응")
    table(doc, ["기준", "해당 내용"], [
        ["NVIDIA Agent 기술 활용 심도", "Nemotron 3 Ultra를 함수 호출과 JSON 스키마 강제, 단계별 reasoning 모드 제어로 사용한다. 판독 에이전트 3종과 임원 에이전트 7인이 병렬로 돌고, 인용 검증과 입장 집계는 코드가 한다. Agent Skills 규격의 스킬과 OpenShell 정책 명세를 포함한다."],
        ["실용성, 산업가치, 혁신성", "의학부의 실제 업무 흐름과 규제 제약을 그대로 설계에 넣었다. 요약이 아니라 원문 위치가 붙은 근거를 만들고, 상업 액션은 코드가 막는다."],
        ["완성도", "콘솔 7화면이 배포돼 있고 서명, 심의, 판정을 화면에서 실제로 실행할 수 있다. 실행 기록과 캐시로 같은 결과를 재현한다."],
        ["커스터마이징, 독창성", "환자군과 신호 유형을 사람이 정하는 고정 헤더, 원문 위치를 요구하는 인용 검증기, 판단 문장의 5단계 표기(사실, 패턴, 해석, 제안, 실행), 확신도 가중 집계로 정하는 권고."],
    ], [1.8, 4.9])

    # 7. 한계
    heading(doc, "7. 한계와 남은 일")
    bullets(doc, [
        "OpenShell 안에서 실행하지 않았다. 개발 환경이 macOS라 정책 명세만 만들었고, 본선에서 Brev 환경에 적용하는 것이 첫 작업이다.",
        "권역은 가상 의료진마다 하나씩 결정론적으로 부여한 합성 값이라 지도와 격자는 권역 해상도까지만 그린다. 원본 콘솔 화면 중 데이터 계약 편집, 배치 판독, 시뮬레이터, 시장 화면은 이번 예선 백엔드가 제공하지 않아 메뉴에서 뺐다.",
        "같은 근거로 심의를 다시 열면 권고가 보류와 기각 사이에서 달라질 수 있다. 결정은 사람이 하므로 권고는 참고값이다. 같은 서명자와 같은 근거면 캐시로 같은 회의가 재생된다.",
        f"무료 NIM 엔드포인트는 분당 호출 수 제한이 있어 면담 {f['notes']}건 추출에 약 40분이 걸렸고 한 번은 속도 제한으로 중단돼 백오프를 늘려 다시 돌렸다.",
        "근거 후보 선별은 PubMed 관련도순에 무작위 대조시험 우선 조회를 더한 것이다. 임베딩 기반 선별은 하지 않았다.",
    ])

    # 부록
    heading(doc, "부록 A. 실행 방법")
    para(doc, "백엔드: uv sync 뒤 .env에 NVIDIA_API_KEY(build.nvidia.com에서 무료 발급)를 넣고 uv run uvicorn loop.web:app --port 8030. "
              "콘솔: console 폴더에서 npm install 뒤 NEXT_PUBLIC_API_BASE_URL=http://localhost:8030/api npm run dev. "
              "명령줄로도 같은 단계를 실행할 수 있다(uv run python -m loop.cli sense | screen | review | board | approve). "
              "공개 API 응답과 모델 응답이 저장소에 캐시돼 있어 키 없이도 같은 입력은 즉시 재생된다.")
    heading(doc, "부록 B. 데이터와 보안")
    para(doc, "면담 기록은 전부 합성이며 실제 환자, 의료진, 기관 정보가 없다. 근거는 공개 API만 쓴다. API 키는 호스트 환경변수에만 있고 캐시, 로그, 저장소에 남지 않는다. "
              "무료 엔드포인트 약관상 입력이 서비스 개선에 쓰일 수 있으므로 합성 데이터만 보냈다.")
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"→ {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    build()
